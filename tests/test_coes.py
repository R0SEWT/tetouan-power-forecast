import datetime as dt
import json

import openpyxl
import pandas as pd
import pytest

from tetouan_power_forecast.coes import (
    BASE_IEOD,
    anexos,
    bronze_ieod,
    carpeta_del_dia,
    carpeta_del_mes,
    leer_demanda_areas,
    parsear_listado,
    silver_areas,
    un_anexo_por_anio,
)


def _horas_2017():
    """Como el anexo de 2017: horas de Excel y las 24:00 como 1900-01-01."""
    horas = [dt.time(h, m) for h in range(24) for m in (0, 30)][1:]
    return horas + [dt.datetime(1900, 1, 1, 0, 0)]


def _horas_2026(ultima="24:00"):
    """Como el anexo de 2026: horas como texto."""
    return [f"{h:02d}:{m:02d}" for h in range(24) for m in (0, 30)][1:] + [ultima]


def _anexo(ruta, formato="2017", base=6000.0, ultima="24:00", con_hoja=True):
    libro = openpyxl.Workbook()
    libro.active.title = "RESUMEN"
    if con_hoja:
        hoja = libro.create_sheet("DEMANDA_AREAS")
        hoja.append(["REPORTE DE LA DEMANDA"])
        if formato == "2017":  # la tabla empieza en la columna B, con la hora en B
            hoja.append([None, "FECHAHORA", "DEMANDA", "AREA", "SUBAREAS"])
            hoja.append([None, None, "IMPORTACIÓN ECUADOR", "SEIN", "NORTE", "CENTRO ", "SUR", "LIMA"])
            for i, hora in enumerate(_horas_2017()):
                n, c, s = 1000.0 + i, 3500.0 + i, base - 4500.0
                hoja.append([None, hora, 0, n + c + s, n, c, s, 1500.0])
            hoja.append([None, "MÁXIMO", None, 9999.0, 1.0, 1.0, 1.0])
        else:  # 2026: la hora como texto en la columna A y los números como texto
            hoja.append(["FECHA HORA", "DEMANDA", "AREA", "SUBAREAS"])
            hoja.append([None, "INTERCAMBIOS INTERNACIONALES", "SEIN", "NORTE", "CENTRO", "SUR", "ELECTROANDES"])
            for i, hora in enumerate(_horas_2026(ultima)):
                n, c, s = 1100.0 + i, 4100.0 + i, base - 5200.0
                hoja.append([hora, "0", f"{n + c + s:.2f}", f"{n:.2f}", f"{c:.2f}", f"{s:.2f}", "280"])
            hoja.append(["ENERGÍA", None, "150000"])
    libro.save(ruta)
    return ruta


def test_lee_el_formato_de_2017(tmp_path):
    df = leer_demanda_areas(_anexo(tmp_path / "Anexo_0203.xlsx", "2017"), dt.date(2017, 3, 2))
    assert len(df) == 48
    assert df["DateTime"].iloc[0] == pd.Timestamp("2017-03-02 00:30")
    assert df["DateTime"].iloc[-1] == pd.Timestamp("2017-03-03 00:00")  # las 24:00
    assert (df[["NORTE", "CENTRO", "SUR"]].sum(axis=1) == df["SEIN"]).all()
    assert df["SEIN"].max() < 9999  # la fila de máximo al pie no entra


@pytest.mark.parametrize("ultima", ["24:00", "00:00"])
def test_lee_el_formato_de_2026_con_numeros_como_texto(tmp_path, ultima):
    df = leer_demanda_areas(_anexo(tmp_path / "AnexoA_2210.xlsx", "2026", ultima=ultima), dt.date(2026, 10, 22))
    assert len(df) == 48 and df["DateTime"].is_unique
    assert df["DateTime"].iloc[-1] == pd.Timestamp("2026-10-23 00:00")
    assert df["NORTE"].iloc[0] == 1100.0
    assert df["SEIN"].dtype == float


def test_un_anexo_sin_la_hoja_devuelve_none(tmp_path):
    assert leer_demanda_areas(_anexo(tmp_path / "Anexo.xlsx", con_hoja=False), dt.date(2020, 1, 1)) is None


def test_listado_y_carpetas_con_las_variantes_de_cada_anio():
    texto = ("<li><a onclick=\"openBlob('Post Operaci&#243;n/Reportes/IEOD/2017/03_MARZO/', 'D', '')\">"
             "<a onclick=\"openBlob('Post Operaci&#243;n/Reportes/IEOD/2017/03_MARZO/', 'D', '')\">")
    assert parsear_listado(texto) == [("D", f"{BASE_IEOD}2017/03_MARZO/")]
    meses = [("D", f"{BASE_IEOD}2023/02_Febrero/"), ("D", f"{BASE_IEOD}2023/03_Marzo/"),
             ("D", f"{BASE_IEOD}2023/12 DICIEMBRE/")]
    assert carpeta_del_mes(meses, 3).endswith("03_Marzo/")
    assert carpeta_del_mes(meses, 12).endswith("12 DICIEMBRE/")
    dias = [("D", "x/Dia 01/"), ("D", "x/Día 02/"), ("D", "x/15/")]
    assert carpeta_del_dia(dias, 2) == "x/Día 02/" and carpeta_del_dia(dias, 15) == "x/15/"
    archivos = [("F", "d/RSF_0203.xlsx"), ("F", "d/Anexo_0203.xlsx"), ("F", "d/Anexo1_Resumen_1503.xlsx"),
                ("F", "d/AnexoA.pdf")]
    assert anexos(archivos) == ["d/Anexo_0203.xlsx", "d/Anexo1_Resumen_1503.xlsx"]


class _PortalFalso:
    """Un IEOD de mentira: días con anexo en formato 2026, salvo los de `sin_anexo`."""

    def __init__(self, tmp_path, sin_anexo=()):
        self.tmp_path, self.sin_anexo, self.descargas = tmp_path, set(sin_anexo), []

    def listar(self, ruta):
        partes = ruta[len(BASE_IEOD):].strip("/").split("/")
        if len(partes) == 1:  # el año: sus meses
            return [("D", f"{ruta}{m:02d}_Mes/") for m in range(1, 13)]
        if len(partes) == 2:  # el mes: sus días
            return [("D", f"{ruta}Día {d:02d}/") for d in range(1, 32)]
        anio, mes, dia = int(partes[0]), int(partes[1][:2]), int(partes[2][-2:])
        fecha = dt.date(anio, mes, dia)
        archivos = [("F", f"{ruta}RSF.xlsx")]
        if fecha not in self.sin_anexo:
            archivos.append(("F", f"{ruta}AnexoA_{fecha:%d%m}.xlsx"))
        return archivos

    def descargar(self, ruta):
        self.descargas.append(ruta)
        return _anexo(self.tmp_path / "plantilla.xlsx", "2026").read_bytes()


def test_bronze_retoma_sin_repetir_y_marca_los_dias_sin_anexo(tmp_path):
    portal = _PortalFalso(tmp_path, sin_anexo={dt.date(2026, 1, 2)})
    raiz = tmp_path / "datos"
    bronze_ieod(raiz, desde=dt.date(2026, 1, 1), hasta=dt.date(2026, 1, 3), listar=portal.listar,
                descargar=portal.descargar, hilos=2, progreso=lambda texto: None)
    lineas = [json.loads(x) for x in (raiz / "coes/bronze/ieod/2026/procedencia.jsonl").read_text().splitlines()]
    assert [x["estado"] for x in lineas] == ["ok", "sin_anexo", "ok"]
    assert lineas[0]["archivos"][0]["hojas"] == ["RESUMEN", "DEMANDA_AREAS"]
    assert (raiz / "coes/bronze/ieod/2026/2026-01-01/AnexoA_0101.xlsx").exists()
    assert len(portal.descargas) == 2

    def falla(ruta):
        raise AssertionError("no debía volver a descargar")

    bronze_ieod(raiz, desde=dt.date(2026, 1, 1), hasta=dt.date(2026, 1, 3), listar=portal.listar, descargar=falla,
                progreso=lambda texto: None)


def test_un_dia_sin_anexo_se_vuelve_a_buscar_en_la_corrida_siguiente(tmp_path):
    raiz = tmp_path / "datos"
    portal = _PortalFalso(tmp_path, sin_anexo={dt.date(2026, 1, 2)})
    bronze_ieod(raiz, desde=dt.date(2026, 1, 1), hasta=dt.date(2026, 1, 2), listar=portal.listar,
                descargar=portal.descargar, progreso=lambda texto: None)
    portal.sin_anexo = set()  # el COES lo publicó después
    bronze_ieod(raiz, desde=dt.date(2026, 1, 1), hasta=dt.date(2026, 1, 2), listar=portal.listar,
                descargar=portal.descargar, progreso=lambda texto: None)
    assert portal.descargas[-1].endswith("AnexoA_0201.xlsx") and len(portal.descargas) == 2
    assert (raiz / "coes/bronze/ieod/2026/2026-01-02/AnexoA_0201.xlsx").exists()


def test_silver_continuo_cada_30_minutos_con_auditoria(tmp_path):
    portal = _PortalFalso(tmp_path, sin_anexo={dt.date(2026, 1, 2)})
    raiz = tmp_path / "datos"
    bronze_ieod(raiz, desde=dt.date(2026, 1, 1), hasta=dt.date(2026, 1, 3), listar=portal.listar,
                descargar=portal.descargar, progreso=lambda texto: None)
    df = pd.read_parquet(silver_areas(raiz, progreso=lambda texto: None))
    assert list(df["DateTime"].diff().dropna().unique()) == [pd.Timedelta("30min")]
    assert df["DateTime"].min() == pd.Timestamp("2026-01-01 00:30")
    assert df["SEIN"].isna().sum() == 48  # el día sin anexo: del 2-ene 00:30 a sus 24:00 (3-ene 00:00)
    auditoria = (raiz / "coes/silver/auditoria.md").read_text()
    assert "| Norte + Centro + Sur = SEIN (±1 MW) | ok |" in auditoria
    assert "2026-01-02" in auditoria and "REVISAR" in auditoria


def test_punto_de_control_un_anexo_por_anio(tmp_path):
    portal = _PortalFalso(tmp_path)
    tabla = un_anexo_por_anio(tmp_path / "datos", anios=[2017, 2026], listar=portal.listar,
                              descargar=portal.descargar, progreso=lambda texto: None)
    assert list(tabla["periodos"]) == [48, 48] and (tabla["desvio_max_mw"] < 0.01).all()
    un_anexo_por_anio(tmp_path / "datos", anios=[2017], listar=portal.listar, descargar=lambda r: 1 / 0,
                      progreso=lambda texto: None)  # ya está en la procedencia: no descarga de nuevo


def test_horas_repetidas_en_el_anexo_fallan(tmp_path):
    ruta = _anexo(tmp_path / "Anexo.xlsx", "2026")
    libro = openpyxl.load_workbook(ruta)
    hoja = libro["DEMANDA_AREAS"]
    hoja.cell(row=5, column=1, value="00:30")  # la segunda fila de datos repite las 00:30
    libro.save(ruta)
    with pytest.raises(ValueError, match="horas repetidas"):
        leer_demanda_areas(ruta, dt.date(2026, 1, 1))
