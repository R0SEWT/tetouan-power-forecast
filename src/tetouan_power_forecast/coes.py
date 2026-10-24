"""Demanda del SEIN por área operativa (Norte, Centro, Sur) desde el IEOD del COES.

El COES publica cada día el Informe de Evaluación de la Operación Diaria (IEOD) con un anexo en Excel. Su hoja
`DEMANDA_AREAS` trae la demanda cada 30 minutos del SEIN y de sus tres áreas. Las carpetas, el nombre del anexo y
la forma de la hoja cambian según el año (`03_MARZO` o `03_Marzo`, `Dia 01` o `Día 01`, `Anexo_0203.xlsx` o
`AnexoA_2210.xlsx`; la hora como texto o como hora de Excel), así que todo se busca por contenido.

Capas, dentro de la carpeta de datos:

- coes/bronze/ieod/AAAA/AAAA-MM-DD/: los anexos del día sin tocar; y coes/bronze/ieod/AAAA/procedencia.jsonl,
  una línea por día con la ruta en el COES, el sha256, las hojas que trae cada archivo y la fecha de descarga.
- coes/silver/demanda_areas.parquet y auditoria.md: la demanda cada 30 minutos, con la marca de tiempo al final
  de cada periodo, como la publica el COES (00:30 es el periodo de 00:00 a 00:30).

Las descargas son pesadas (~400 MB por año) y se corren en Colab, no en local.
"""

import datetime as dt
import hashlib
import html
import json
import re
import time
import urllib.error
import urllib.parse
import urllib.request
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

from tetouan_power_forecast.medallon import _reemplazar, _ruta

URL_PORTAL = "https://www.coes.org.pe/Portal"
BASE_IEOD = "Post Operación/Reportes/IEOD/"
AREAS = ["SEIN", "NORTE", "CENTRO", "SUR"]
INICIO = dt.date(2015, 1, 1)
CABECERAS = {"User-Agent": "Mozilla/5.0 (X11; Linux x86_64) tetouan-power-forecast (investigación académica)"}
FUENTE = {
    "fuente": "COES SINAC, Informe de Evaluación de la Operación Diaria (IEOD), anexo diario",
    "pagina": f"{URL_PORTAL}/browser",
    "uso": "Datos públicos del operador del SEIN, sin licencia publicada; se cita al COES como fuente.",
}


def _pedir(url, datos=None, timeout=120, intentos=4, espera=10):
    """GET o POST con reintentos; el portal del COES a veces corta la conexión."""
    cuerpo = urllib.parse.urlencode(datos).encode() if datos is not None else None
    peticion = urllib.request.Request(url, data=cuerpo, headers=CABECERAS)
    for intento in range(1, intentos + 1):
        try:
            with urllib.request.urlopen(peticion, timeout=timeout) as respuesta:
                return respuesta.read()
        except (urllib.error.URLError, TimeoutError, ConnectionError):
            if intento == intentos:
                raise
            time.sleep(espera * intento)


def parsear_listado(texto):
    """Carpetas ('D') y archivos ('F') de una página del navegador de archivos, sin repetir y en orden."""
    vistos = {}
    for ruta, tipo in re.findall(r"openBlob\('([^']+)',\s*'(\w)'", texto):
        vistos.setdefault(html.unescape(ruta), tipo)
    return [(tipo, ruta) for ruta, tipo in vistos.items()]


def listar(ruta):
    """Lista una carpeta del IEOD en el portal (ruta relativa a la raíz del navegador)."""
    datos = {"baseDirectory": BASE_IEOD, "url": ruta, "indicador": "S"}
    return parsear_listado(_pedir(f"{URL_PORTAL}/browser/vistadatos", datos).decode("utf-8", "replace"))


def _nombre(ruta):
    return ruta.rstrip("/").rsplit("/", 1)[-1]


def carpeta_del_mes(listado_anio, mes):
    """La carpeta del mes: empieza con su número de dos cifras (`03_MARZO`, `03_Marzo`, `03 MARZO`)."""
    for tipo, ruta in listado_anio:
        if tipo == "D" and re.match(rf"^{mes:02d}\D", _nombre(ruta)):
            return ruta
    return None


def carpeta_del_dia(listado_mes, dia):
    """La carpeta del día: termina con su número (`Dia 01`, `Día 01`, `01`)."""
    for tipo, ruta in listado_mes:
        numero = re.search(r"(\d{1,2})$", _nombre(ruta))
        if tipo == "D" and numero and int(numero.group(1)) == dia:
            return ruta
    return None


def anexos(listado_dia):
    """Los anexos en Excel del día (`Anexo_0203.xlsx`, `Anexo1_Resumen_1503.xlsx`, `AnexoA_2210.xlsx`)."""
    return [ruta for tipo, ruta in listado_dia if tipo == "F" and re.match(r"(?i)^anexo.*\.xlsx$", _nombre(ruta))]


class Navegador:
    """Encuentra los anexos de cada día, con las carpetas de años y meses en memoria para no repetir consultas."""

    def __init__(self, listar=listar):
        self._listar = listar
        self._cache = {}

    def _listado(self, ruta):
        if ruta not in self._cache:
            self._cache[ruta] = self._listar(ruta)
        return self._cache[ruta]

    def anexos_del_dia(self, dia):
        anio = f"{BASE_IEOD}{dia.year}/"
        mes = carpeta_del_mes(self._listado(anio), dia.month)
        if mes is None:
            return []
        carpeta = carpeta_del_dia(self._listado(mes), dia.day)
        return anexos(self._listar(carpeta)) if carpeta else []


def _descargar_archivo(ruta):
    return _pedir(f"{URL_PORTAL}/browser/download?" + urllib.parse.urlencode({"url": ruta}), timeout=300)


def _hojas(contenido_o_ruta):
    libro = openpyxl.load_workbook(contenido_o_ruta, read_only=True)
    try:
        return libro.sheetnames
    finally:
        libro.close()


def _dias(desde, hasta):
    return [desde + dt.timedelta(days=k) for k in range((hasta - desde).days + 1)]


def _procedencia(raiz, anio):
    return _ruta(raiz, f"coes/bronze/ieod/{anio}", "procedencia.jsonl")


def dias_bajados(raiz, anio):
    """Fechas del año que ya tienen una línea en la procedencia (con anexo o marcadas sin anexo)."""
    archivo = _procedencia(raiz, anio)
    if not archivo.exists():
        return {}
    lineas = (json.loads(linea) for linea in archivo.read_text().splitlines() if linea.strip())
    return {linea["fecha"]: linea for linea in lineas}


def _bajar_dia(raiz, dia, navegador, descargar):
    carpeta = _ruta(raiz, f"coes/bronze/ieod/{dia.year}/{dia:%Y-%m-%d}", "x").parent
    archivos = []
    for ruta in navegador.anexos_del_dia(dia):
        contenido = descargar(ruta)
        destino = carpeta / _nombre(ruta)
        _reemplazar(destino, lambda r, c=contenido: r.write_bytes(c))
        archivos.append({"ruta_coes": ruta, "archivo": destino.name, "bytes": len(contenido),
                         "sha256": hashlib.sha256(contenido).hexdigest(), "hojas": _hojas(destino)})
    return {"fecha": f"{dia:%Y-%m-%d}", "estado": "ok" if archivos else "sin_anexo", "archivos": archivos,
            "descargado_utc": dt.datetime.now(dt.UTC).isoformat(timespec="seconds")}


def bronze_ieod(raiz=None, desde=INICIO, hasta=None, listar=listar, descargar=_descargar_archivo, hilos=4,
                progreso=print):
    """Baja los anexos del IEOD día por día, sin tocarlos. Se puede cortar y retomar.

    Un día ya registrado en la procedencia no se vuelve a pedir; los días sin anexo quedan marcados como
    `sin_anexo` para revisarlos, no se reintentan solos. Los días de un mes se bajan con `hilos` en paralelo.
    """
    hasta = hasta or (dt.datetime.now(dt.UTC).date() - dt.timedelta(days=2))  # el IEOD sale en D+1
    navegador = Navegador(listar)
    for anio in range(desde.year, hasta.year + 1):
        hechos = dias_bajados(raiz, anio)
        pendientes = [d for d in _dias(max(desde, dt.date(anio, 1, 1)), min(hasta, dt.date(anio, 12, 31)))
                      if f"{d:%Y-%m-%d}" not in hechos]
        for mes in sorted({(d.year, d.month) for d in pendientes}):
            dias_mes = [d for d in pendientes if (d.year, d.month) == mes]
            with ThreadPoolExecutor(max_workers=hilos) as pool:
                lineas = list(pool.map(lambda d: _bajar_dia(raiz, d, navegador, descargar), dias_mes))
            with _procedencia(raiz, anio).open("a") as salida:  # solo se agrega: lo ya escrito no se pierde
                for linea in sorted(lineas, key=lambda x: x["fecha"]):
                    salida.write(json.dumps(linea, ensure_ascii=False) + "\n")
            sin = sum(linea["estado"] != "ok" for linea in lineas)
            progreso(f"{mes[0]}-{mes[1]:02d}: {len(lineas)} días" + (f", {sin} sin anexo" if sin else ""))
    texto = json.dumps({**FUENTE, "desde": f"{desde}", "hasta": f"{hasta}",
                        "actualizado_utc": dt.datetime.now(dt.UTC).isoformat(timespec="seconds")},
                       ensure_ascii=False, indent=2) + "\n"
    _reemplazar(_ruta(raiz, "coes/bronze/ieod", "procedencia.json"), lambda r: r.write_text(texto))
    return _ruta(raiz, "coes/bronze/ieod", "x").parent


def _hora(valor):
    """(hora, minuto) de una celda de hora del IEOD, o None si la celda no es una hora."""
    if isinstance(valor, dt.datetime):
        return valor.hour, valor.minute
    if isinstance(valor, dt.time):
        return valor.hour, valor.minute
    if isinstance(valor, str):
        hora = re.match(r"^\s*(\d{1,2}):(\d{2})", valor)
        if hora:
            return int(hora.group(1)), int(hora.group(2))
    return None


def _numero(valor):
    try:
        return float(valor)
    except (TypeError, ValueError):
        return np.nan


def leer_demanda_areas(archivo, dia):
    """La demanda del SEIN y de sus tres áreas en un anexo, o None si el anexo no trae la hoja.

    La marca de tiempo es el final de cada periodo de 30 minutos; la última del día (24:00, que el COES escribe
    como 00:00 o como 1900-01-01) pasa a las 00:00 del día siguiente.
    """
    libro = openpyxl.load_workbook(archivo, read_only=True, data_only=True)
    try:
        if "DEMANDA_AREAS" not in libro.sheetnames:
            return None
        filas = list(libro["DEMANDA_AREAS"].iter_rows(values_only=True))
    finally:
        libro.close()
    normal = [[str(c).strip().upper() if isinstance(c, str) else None for c in fila] for fila in filas]
    i_cab = next((i for i, fila in enumerate(normal) if "SEIN" in fila), None)
    if i_cab is None:
        raise ValueError(f"{archivo}: la hoja DEMANDA_AREAS no tiene la cabecera SEIN")
    columnas = {area: normal[i_cab].index(area) for area in AREAS if area in normal[i_cab]}
    faltan = set(AREAS) - set(columnas)
    if faltan:
        raise ValueError(f"{archivo}: faltan las áreas {sorted(faltan)} en DEMANDA_AREAS")
    registros = []
    for fila in filas[i_cab + 1:]:
        hora = next((_hora(c) for c in fila[:columnas["SEIN"]] if _hora(c) is not None), None)
        if hora is None:
            continue  # filas de máximo, mínimo, promedio o energía al pie
        minutos = hora[0] * 60 + hora[1]
        if minutos == 0:
            minutos = 24 * 60
        fecha = pd.Timestamp(dia) + pd.Timedelta(minutes=minutos)
        registros.append({"DateTime": fecha, **{a: _numero(fila[c]) for a, c in columnas.items()}})
    return pd.DataFrame(registros, columns=["DateTime", *AREAS])


def auditar_areas(df, sin_anexo=()):
    esperado = pd.date_range(df["DateTime"].min(), df["DateTime"].max(), freq="30min")
    suma = df[["NORTE", "CENTRO", "SUR"]].sum(axis=1, min_count=3)
    desvio = (suma - df["SEIN"]).abs()
    faltantes = int(df[AREAS].isna().any(axis=1).sum())
    chequeos = [
        ("Un registro cada 30 minutos", len(df) == len(esperado), f"{len(df):,} de {len(esperado):,}"),
        ("Sin fechas repetidas", not df["DateTime"].duplicated().any(), f"{int(df['DateTime'].duplicated().sum())}"),
        ("Norte + Centro + Sur = SEIN (±1 MW)", bool((desvio.dropna() <= 1).all()),
         f"{int((desvio > 1).sum()):,} periodos fuera; desvío máximo {desvio.max():.1f} MW"),
        ("Sin faltantes", faltantes == 0, f"{faltantes:,} periodos con algún área vacía ({faltantes / len(df):.2%})"),
        ("Demanda positiva", bool((df[AREAS].dropna() > 0).all().all()),
         f"{int((df[AREAS] <= 0).sum().sum())} valores no positivos"),
        ("Días con anexo", not sin_anexo, f"{len(sin_anexo)} días sin anexo: {', '.join(sorted(sin_anexo)[:10])}"),
    ]
    return pd.DataFrame(chequeos, columns=["chequeo", "ok", "detalle"])


def silver_areas(raiz=None, forzar=False, progreso=print):
    """La demanda por área cada 30 minutos de todo el bronze, con su auditoría. Reporta los huecos, no corta."""
    destino = _ruta(raiz, "coes/silver", "demanda_areas.parquet")
    if destino.exists() and not forzar:
        return destino
    bronze = _ruta(raiz, "coes/bronze/ieod", "x").parent
    partes, sin_anexo = [], []
    for procedencia in sorted(bronze.glob("*/procedencia.jsonl")):
        for linea in dias_bajados(raiz, procedencia.parent.name).values():
            leida = None
            for archivo in linea["archivos"]:
                if "DEMANDA_AREAS" in archivo["hojas"]:
                    leida = leer_demanda_areas(procedencia.parent / linea["fecha"] / archivo["archivo"],
                                               dt.date.fromisoformat(linea["fecha"]))
                    break
            if leida is None:
                sin_anexo.append(linea["fecha"])
            else:
                partes.append(leida)
        progreso(f"{procedencia.parent.name}: leído")
    df = (pd.concat(partes).drop_duplicates("DateTime", keep="last").sort_values("DateTime")
          .set_index("DateTime"))
    df = df.reindex(pd.date_range(df.index.min(), df.index.max(), freq="30min")).rename_axis("DateTime")
    df = df.reset_index()
    auditoria = auditar_areas(df, sin_anexo)
    reporte = ["# Auditoría del silver de demanda por área (IEOD del COES)", "",
               (f"Registros: {len(df):,} · desde {df['DateTime'].min()} hasta {df['DateTime'].max()} "
                "(hora de Lima, fin de cada periodo)"), "",
               "| Chequeo | Resultado | Detalle |", "|---|---|---|",
               *(f"| {c} | {'ok' if ok else 'REVISAR'} | {d} |" for c, ok, d in auditoria.itertuples(index=False))]
    destino.with_name("auditoria.md").write_text("\n".join(reporte) + "\n")
    _reemplazar(destino, lambda ruta: df.to_parquet(ruta, index=False))
    return destino


def cargar_demanda_areas(raiz=None):
    return pd.read_parquet(silver_areas(raiz))


def un_anexo_por_anio(raiz=None, anios=range(2015, 2027), listar=listar, descargar=_descargar_archivo,
                      progreso=print):
    """Punto de control antes de la descarga masiva: el 2 de marzo de cada año, leído con el mismo lector.

    Devuelve una tabla con el archivo, las hojas y si Norte + Centro + Sur = SEIN en ese día.
    """
    navegador = Navegador(listar)
    filas = []
    for anio in anios:
        dia = dt.date(anio, 3, 2)
        linea = dias_bajados(raiz, anio).get(f"{dia:%Y-%m-%d}")
        if linea is None:  # se registra en la procedencia: la descarga masiva no lo vuelve a pedir
            linea = _bajar_dia(raiz, dia, navegador, descargar)
            with _procedencia(raiz, anio).open("a") as salida:
                salida.write(json.dumps(linea, ensure_ascii=False) + "\n")
        fila = {"anio": anio, "estado": linea["estado"], "archivos": [a["archivo"] for a in linea["archivos"]]}
        for archivo in linea["archivos"]:
            if "DEMANDA_AREAS" in archivo["hojas"]:
                df = leer_demanda_areas(Path(_ruta(raiz, f"coes/bronze/ieod/{anio}/{dia}", archivo["archivo"])), dia)
                desvio = (df[["NORTE", "CENTRO", "SUR"]].sum(axis=1) - df["SEIN"]).abs().max()
                fila.update({"periodos": len(df), "desvio_max_mw": round(float(desvio), 3),
                             "hojas": archivo["hojas"]})
                break
        filas.append(fila)
        progreso(f"{anio}: {fila['estado']} · {fila.get('periodos', 0)} periodos")
    return pd.DataFrame(filas)
